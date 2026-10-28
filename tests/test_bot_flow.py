import asyncio
from unittest.mock import patch

import pytest

from app import bot_handlers
from app.ai_check import BillExtractionResult

from conftest import (
    FakeMessage,
    FakeUpdate,
    FakeUser,
    button_callbacks,
    button_labels,
    make_context,
    make_settings,
    reply_texts,
    session_of,
)


class Driver:
    """Simulates one Telegram user talking to the bot."""

    def __init__(self, context, user=None):
        self.context = context
        self.user = user or FakeUser()
        self.message = FakeMessage()

    async def click(self, data):
        update = FakeUpdate(user=self.user, data=data, message=self.message)
        await bot_handlers.route_update(update, self.context)

    async def send(self, text):
        self.message.text = text
        self.message.photo = None
        self.message.document = None
        update = FakeUpdate(user=self.user, message=self.message)
        await bot_handlers.route_update(update, self.context)

    async def send_photo(self, file_id="photo1", file_size=1024):
        self.message.text = None
        self.message.photo = [
            type("P", (), {"file_id": file_id, "file_size": file_size})()
        ]
        self.message.document = None
        update = FakeUpdate(user=self.user, message=self.message)
        await bot_handlers.route_update(update, self.context)

    def texts(self):
        return reply_texts(self.message)

    def session(self):
        return session_of(self.context, self.user.id)


async def test_reimbursement_question_order_ai_disabled():
    driver = Driver(make_context())
    await driver.click("new")
    assert "reimbursement month" in driver.texts()[-1]

    await driver.click("month:2026-10")
    assert "department" in driver.texts()[-1].lower()

    await driver.click("department:0")
    assert "program" in driver.texts()[-1].lower()

    await driver.click("program:0")
    assert "project" in driver.texts()[-1].lower()

    await driver.click("project:0")
    # Bill upload must come before expense type and before bill fields.
    assert "bill" in driver.texts()[-1].lower()
    assert driver.session()["step"] == "upload"

    await driver.send_photo()
    await driver.click("bills:done")
    assert "expense type" in driver.texts()[-1].lower()
    assert driver.session()["step"] == "expense_type"

    await driver.click("expense:0")
    assert "bill date" in driver.texts()[-1].lower()

    await driver.send("15/10/2026")
    assert "vendor" in driver.texts()[-1].lower()

    await driver.send("Acme Corp")
    assert "amount" in driver.texts()[-1].lower()

    await driver.send("1250")
    last = driver.texts()[-1]
    assert "review" in last.lower()
    assert driver.session()["step"] == "summary"
    # Summary shows flow type, name, and all entered values.
    for expected in (
        "Reimbursement Flow", "Test User", "2026-10", "Education",
        "Mindset Program", "VVCE", "Travel", "2026-10-15", "Acme Corp", "1250",
    ):
        assert expected in last, expected


async def test_full_submission_writes_row():
    driver = Driver(make_context())
    await driver.click("new")
    await driver.click("month:2026-10")
    await driver.click("department:0")
    await driver.click("program:0")
    await driver.click("project:0")
    await driver.send_photo()
    await driver.click("bills:done")
    await driver.click("expense:0")
    await driver.send("15/10/2026")
    await driver.send("Acme Corp")
    await driver.send("1250")
    await driver.click("submit")

    sheets = driver.context.application.bot_data["sheets"]
    assert len(sheets.appended) == 1
    row = sheets.appended[0]
    assert row[5] == "2026-10"           # Month
    assert row[6] == "Education"         # Department
    assert row[7] == "Travel"            # Expense Type
    assert row[9] == "2026-10-15"        # Bill Date
    assert row[10] == "Acme Corp"        # Vendor
    assert row[11] == 1250.0             # Amount
    assert row[12] == "VVCE"             # Project
    assert "https://drive.test/" in row[13]
    assert row[16] == "Pending"          # Status
    assert row[18] == "Reimbursement Flow"
    assert row[19] == "Mindset Program"


async def test_ai_confirm_skips_manual_fields():
    context = make_context(settings=make_settings(ai_check_enabled=True, nvidia_api_key="k"))
    driver = Driver(context)
    extraction = BillExtractionResult(
        "OK", bill_date="2026-10-15", vendor="Acme Corp", amount="1250.00",
        bill_number="INV-123",
    )
    with patch.object(
        bot_handlers, "extract_bill_details", return_value=extraction
    ):
        await driver.click("new")
        await driver.click("month:2026-10")
        await driver.click("department:0")
        await driver.click("program:0")
        await driver.click("project:0")
        await driver.send_photo()
        await driver.click("bills:done")

    assert "detected from your bill" in driver.texts()[-1]
    assert "15/10/2026" in driver.texts()[-1]
    assert "INV-123" in driver.texts()[-1]
    assert "ai:confirm" in button_callbacks(driver.message.replies[-1])

    await driver.click("ai:confirm")
    assert "expense type" in driver.texts()[-1].lower()

    await driver.click("expense:0")
    # All fields confirmed by AI -> straight to summary, no date/vendor/amount prompts.
    assert "review" in driver.texts()[-1].lower()
    session = driver.session()
    assert session["bill_date"] == "2026-10-15"
    assert session["vendor"] == "Acme Corp"
    assert session["amount"] == "1250.00"
    assert session["bill_number"] == "INV-123"


async def test_ai_edit_single_field_preserves_others():
    context = make_context(settings=make_settings(ai_check_enabled=True, nvidia_api_key="k"))
    driver = Driver(context)
    extraction = BillExtractionResult(
        "OK", bill_date="2026-10-15", vendor="Acme Corp", amount="1250.00"
    )
    with patch.object(
        bot_handlers, "extract_bill_details", return_value=extraction
    ):
        await driver.click("new")
        await driver.click("month:2026-10")
        await driver.click("department:0")
        await driver.click("program:0")
        await driver.click("project:0")
        await driver.send_photo()
        await driver.click("bills:done")

    await driver.click("ai:edit:amount")
    assert "amount" in driver.texts()[-1].lower()
    await driver.send("2000")
    assert "expense type" in driver.texts()[-1].lower()
    await driver.click("expense:0")

    session = driver.session()
    assert session["amount"] == "2000.00"
    assert session["bill_date"] == "2026-10-15"   # preserved
    assert session["vendor"] == "Acme Corp"       # preserved
    assert "2000" in driver.texts()[-1]


async def test_ai_unreadable_falls_back_to_manual_entry():
    context = make_context(settings=make_settings(ai_check_enabled=True, nvidia_api_key="k"))
    driver = Driver(context)
    extraction = BillExtractionResult("Unreadable", notes="The AI could not read this bill.")
    with patch.object(
        bot_handlers, "extract_bill_details", return_value=extraction
    ):
        await driver.click("new")
        await driver.click("month:2026-10")
        await driver.click("department:0")
        await driver.click("program:0")
        await driver.click("project:0")
        await driver.send_photo()
        await driver.click("bills:done")

    assert "expense type" in driver.texts()[-1].lower()
    await driver.click("expense:0")
    assert "bill date" in driver.texts()[-1].lower()
    await driver.send("15/10/2026")
    await driver.send("Acme Corp")
    await driver.send("1250")
    assert "review" in driver.texts()[-1].lower()


async def test_stale_button_press_does_not_refire():
    driver = Driver(make_context())
    await driver.click("new")
    await driver.click("month:2026-10")
    assert driver.session()["step"] == "department"

    # Pressing the month button again (stale keyboard) must not re-run the step.
    await driver.click("month:2026-11")
    session = driver.session()
    assert session["step"] == "department"
    assert session["month"] == "2026-10"
    assert "no longer active" in driver.texts()[-1]


async def test_keyboard_removed_after_valid_click():
    driver = Driver(make_context())
    await driver.click("new")
    await driver.click("month:2026-10")
    assert driver.message.markup_removed >= 1


async def test_submit_advances_once_on_double_click():
    driver = Driver(make_context())
    await driver.click("new")
    await driver.click("month:2026-10")
    await driver.click("department:0")
    await driver.click("program:0")
    await driver.click("project:0")
    await driver.send_photo()
    await driver.click("bills:done")
    await driver.click("expense:0")
    await driver.send("15/10/2026")
    await driver.send("Acme Corp")
    await driver.send("1250")

    sheets = driver.context.application.bot_data["sheets"]
    await driver.click("submit")
    await driver.click("submit")   # second press on the stale summary keyboard
    assert len(sheets.appended) == 1
    assert "already submitted" in driver.texts()[-1]


async def test_concurrent_submits_do_not_duplicate():
    driver = Driver(make_context())
    await driver.click("new")
    await driver.click("month:2026-10")
    await driver.click("department:0")
    await driver.click("program:0")
    await driver.click("project:0")
    await driver.send_photo()
    await driver.click("bills:done")
    await driver.click("expense:0")
    await driver.send("15/10/2026")
    await driver.send("Acme Corp")
    await driver.send("1250")

    # Two callback updates racing each other.
    await asyncio.gather(driver.click("submit"), driver.click("submit"))
    sheets = driver.context.application.bot_data["sheets"]
    assert len(sheets.appended) == 1


async def test_config_driven_option_buttons():
    context = make_context()
    driver = Driver(context)
    await driver.click("new")
    await driver.click("month:2026-10")
    labels = button_labels(driver.message.replies[-1])
    assert labels[:2] == ["Education", "HR"]

    await driver.click("department:0")
    labels = button_labels(driver.message.replies[-1])
    assert "Mindset Program" in labels
    assert "Onboarding" not in labels   # programs are scoped to the department

    await driver.click("program:0")
    labels = button_labels(driver.message.replies[-1])
    assert "VVCE" in labels and "SKIT" in labels


async def test_removed_option_reprompts_at_submit():
    context = make_context()
    driver = Driver(context)
    await driver.click("new")
    await driver.click("month:2026-10")
    await driver.click("department:0")
    await driver.click("program:0")
    await driver.click("project:0")
    await driver.send_photo()
    await driver.click("bills:done")
    await driver.click("expense:0")
    await driver.send("15/10/2026")
    await driver.send("Acme Corp")
    await driver.send("1250")

    # Admin removes the picked project while the draft is at summary.
    sheets = driver.context.application.bot_data["sheets"]
    sheets.config["projects"]["Mindset Program"] = ["VVCE"]

    # project was VVCE (index 0) -> still valid; remove it instead.
    sheets.config["projects"]["Mindset Program"] = ["SKIT"]
    await driver.click("submit")
    assert any("no longer available" in t for t in driver.texts())
    assert driver.session()["step"] == "project"
    assert sheets.appended == []


async def test_no_manual_entry_options_in_menus():
    context = make_context()
    driver = Driver(context)
    await driver.click("new")
    await driver.click("month:2026-10")
    await driver.click("department:0")
    await driver.click("program:0")

    # Collect every callback from the department, program, project, and
    # expense menus; none may offer free-text entry.
    seen = []
    for reply in driver.message.replies:
        seen.extend(button_callbacks(reply))
    await driver.click("project:0")
    await driver.send_photo()
    await driver.click("bills:done")
    for reply in driver.message.replies:
        seen.extend(button_callbacks(reply))
    config_scopes = {"department", "program", "project", "expense"}
    assert not any(
        cb.endswith(":other") and cb.split(":")[0] in config_scopes for cb in seen
    )


async def test_direct_payment_admin_only():
    # Non-admin pressing the direct-payment button is refused.
    context = make_context()
    driver = Driver(context, user=FakeUser(user_id=111))
    await driver.click("flow:direct")
    assert "admins only" in driver.texts()[-1]

    # Admin sees the flow choice and can pick direct payment.
    admin = Driver(context, user=FakeUser(user_id=999))
    await admin.click("new")
    callbacks = button_callbacks(admin.message.replies[-1])
    assert "flow:reimbursement" in callbacks and "flow:direct" in callbacks
    await admin.click("flow:direct")
    assert admin.session()["flow_type"] == "Direct Payment Flow"
    assert "expense month" in admin.texts()[-1]


async def test_cancel_button_works_from_upload():
    driver = Driver(make_context())
    await driver.click("new")
    await driver.click("month:2026-10")
    await driver.click("department:0")
    await driver.click("program:0")
    await driver.click("project:0")
    assert "cancel" in button_callbacks(driver.message.replies[-1])
    await driver.click("cancel")
    assert "cancelled" in driver.texts()[-1]
    assert driver.user.id not in driver.context.application.bot_data["sessions"]


async def test_back_button_navigates_previous_step():
    driver = Driver(make_context())
    await driver.click("new")
    await driver.click("month:2026-10")
    assert driver.session()["step"] == "department"

    await driver.click("back")
    assert driver.session()["step"] == "month"
    assert "reimbursement month" in driver.texts()[-1]

    await driver.click("month:2026-10")
    await driver.click("department:0")
    assert driver.session()["step"] == "program"

    await driver.click("back")
    assert driver.session()["step"] == "department"

    await driver.click("department:0")
    await driver.click("program:0")
    await driver.click("project:0")
    assert driver.session()["step"] == "upload"

    await driver.click("back")
    assert driver.session()["step"] == "project"


async def test_report_generates_docx_and_saves_to_drive():
    from docx import Document
    import io

    context = make_context()
    sheets = context.application.bot_data["sheets"]
    sheets.report_rows = [
        ["2026-10-01T10:00:00+00:00", "id1", "t@x.com", "Test User", "111",
         "2026-10", "Education", "Travel", "", "2026-10-15", "Acme Corp",
         "1250.5", "VVCE", "https://l", "OK", "", "Pending", "",
         "Reimbursement Flow", "Mindset Program"],
        ["2026-10-02T10:00:00+00:00", "id2", "t@x.com", "Test User", "111",
         "2026-10", "Education", "Purchases", "", "2026-10-16", "Bolt Store",
         "300", "VVCE", "https://l", "OK", "", "Pending", "",
         "Reimbursement Flow", "Mindset Program"],
    ]
    driver = Driver(context)
    driver.message.text = "/report"
    update = FakeUpdate(user=driver.user, message=driver.message)
    await bot_handlers.route_update(update, context)

    assert "reimbursement report" in driver.texts()[-1].lower()
    await driver.click("report_month:2026-10")

    bot = context.bot
    assert len(bot.documents) == 1
    _, docx_bytes, filename, caption = bot.documents[0]
    assert filename.startswith("reimbursement_2026-10")
    assert docx_bytes[:2] == b"PK"

    document = Document(io.BytesIO(docx_bytes))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Test User" in text
    assert "Finance Officer" in text
    table_text = "\n".join(
        cell.text for row in document.tables[0].rows for cell in row.cells
    )
    assert "Travel" in table_text and "Purchases" in table_text
    assert "Acme Corp" in table_text
    assert "1,550.50" in table_text  # total

    drive = context.application.bot_data["drive"]
    assert any(u[2].endswith(".docx") for u in drive.uploads)


async def test_reuse_last_details_skips_option_steps():
    context = make_context()
    context.application.bot_data["last_details"] = {
        111: {
            "department": "Education",
            "program": "Mindset Program",
            "project": "VVCE",
            "department_from_options": True,
            "program_from_options": True,
            "project_from_options": True,
        }
    }
    driver = Driver(context)
    await driver.click("new")
    assert "same details" in driver.texts()[-1]
    assert driver.session()["step"] == "reuse"

    await driver.click("reuse:yes")
    assert driver.session()["step"] == "month"

    await driver.click("month:2026-10")
    # Department/program/project reused -> straight to bill upload.
    assert driver.session()["step"] == "upload"
    session = driver.session()
    assert session["department"] == "Education"
    assert session["program"] == "Mindset Program"
    assert session["project"] == "VVCE"


async def test_reuse_no_starts_normal_flow():
    context = make_context()
    context.application.bot_data["last_details"] = {
        111: {
            "department": "Education",
            "program": "Mindset Program",
            "project": "VVCE",
        }
    }
    driver = Driver(context)
    await driver.click("new")
    await driver.click("reuse:no")
    await driver.click("month:2026-10")
    assert driver.session()["step"] == "department"


async def test_stats_menu_shows_month_totals():
    context = make_context()
    month = "2026-10"
    context.application.bot_data["sheets"].report_rows = [
        ["ts", "id1", "e", "Test User", "111", month, "Education", "Travel", "",
         "2026-10-15", "Acme", "1000", "P", "l", "OK", "", "Pending", "", "Reimbursement Flow", "prog"],
        ["ts", "id2", "e", "Test User", "111", month, "Education", "Purchases", "",
         "2026-10-16", "B", "250.50", "P", "l", "OK", "", "Paid", "", "Reimbursement Flow", "prog"],
    ]
    driver = Driver(context)
    await driver.click("menu:stats")
    text = driver.texts()[-1]
    assert "2 expenses" in text
    assert "1,250.50" in text
    assert "Pending: 1" in text and "Paid: 1" in text


async def test_bills_done_requires_a_file():
    driver = Driver(make_context())
    await driver.click("new")
    await driver.click("month:2026-10")
    await driver.click("department:0")
    await driver.click("program:0")
    await driver.click("project:0")
    await driver.click("bills:done")
    assert "at least one bill" in driver.texts()[-1]
    assert driver.session()["step"] == "upload"
