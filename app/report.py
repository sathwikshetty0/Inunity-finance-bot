from __future__ import annotations

import io
from collections import OrderedDict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

HEADERS = ["Sr No", "Category of Expenses", "Bill number", "Purpose of Expenses", "Amount (₹)"]


def _amount(row: list[Any]) -> Decimal:
    try:
        return Decimal(str(row[11]).replace(",", ""))
    except (InvalidOperation, IndexError):
        return Decimal("0")


def _text(row: list[Any], index: int) -> str:
    return str(row[index]).strip() if len(row) > index else ""


def build_reimbursement_docx(
    name: str, month: str, rows: list[list[str]], today: datetime | None = None
) -> bytes:
    """Build the reimbursement-request letter as a .docx and return its bytes.

    `rows` are raw Data-tab rows for one employee and one month.
    """
    today = today or datetime.now()
    doc = Document()
    for section in doc.sections:
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    header = doc.add_paragraph()
    header.add_run("From,")
    header.add_run("\t\t\t\t\t")
    header.add_run(f"Date: {today.strftime('%d/%m/%Y')}")
    doc.add_paragraph(name)
    doc.add_paragraph("InUnity Private Limited")
    doc.add_paragraph("")
    doc.add_paragraph("To,")
    doc.add_paragraph("Finance Officer")
    doc.add_paragraph("InUnity Private Limited")
    doc.add_paragraph("")
    subject = doc.add_paragraph()
    subject.add_run(f"Sub: Request for reimbursement of expenditure for {month}").bold = True
    doc.add_paragraph("")
    doc.add_paragraph("Respected Sir/Madam,")
    doc.add_paragraph(
        "Here's the list of services which I would like to be reimbursed for."
    )

    # Group the submissions by expense type, preserving first-seen order.
    groups: "OrderedDict[str, list[list[str]]]" = OrderedDict()
    for row in rows:
        groups.setdefault(_text(row, 7) or "Other", []).append(row)

    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells, HEADERS):
        cell.text = heading
        for run in cell.paragraphs[0].runs:
            run.bold = True

    total = Decimal("0")
    for serial, (category, items) in enumerate(groups.items(), start=1):
        for index, row in enumerate(items):
            cells = table.add_row().cells
            cells[0].text = str(serial) if index == 0 else ""
            cells[1].text = category if index == 0 else ""
            cells[2].text = _text(row, 8) or "NA"  # Description stores the bill number
            vendor = _text(row, 10) or "NA"
            project = _text(row, 12)
            bill_date = _text(row, 9)
            purpose_parts = [vendor]
            if project:
                purpose_parts.append(project)
            if bill_date:
                purpose_parts.append(bill_date)
            cells[3].text = " - ".join(purpose_parts)
            amount = _amount(row)
            total += amount
            cells[4].text = f"{amount:,.2f}"

    total_cells = table.add_row().cells
    total_cells[3].text = "Total"
    total_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in total_cells[3].paragraphs[0].runs:
        run.bold = True
    total_cells[4].text = f"₹{total:,.2f}"
    for run in total_cells[4].paragraphs[0].runs:
        run.bold = True

    doc.add_paragraph("")
    doc.add_paragraph(
        "I kindly request you to reimburse this amount to me as I have completed "
        "procuring the listed services for InUnity Private Limited."
    )
    doc.add_paragraph("")
    doc.add_paragraph("Sincerely,")
    doc.add_paragraph("")
    doc.add_paragraph("Signature (Employee)")
    doc.add_paragraph("")
    doc.add_paragraph("")
    approved = doc.add_paragraph()
    approved.add_run("Approved By").bold = True
    doc.add_paragraph("Vertical Manager")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
